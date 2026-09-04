from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reddit_scraper.storage.excel_media_embed import embed_media_sheet, embed_post_sheet, prepare_posts_for_media, split_media_rows

try:  # openpyxl >= 3.1
    from openpyxl.worksheet.properties import Outline
except ImportError:  # pragma: no cover - older openpyxl
    from openpyxl.worksheet.properties import OutlineProperties as Outline  # type: ignore[attr-defined]

# (internal key, 中文列名) — 内部字段键保持不变，只有表头展示为中文
POST_FIELDS = [
    ("post_id", "帖子ID"),
    ("subreddit", "子版块"),
    ("subreddit_id", "子版块ID"),
    ("title", "标题"),
    ("author", "作者"),
    ("author_id", "作者ID"),
    ("created", "发布时间"),
    ("score", "点赞数"),
    ("upvote_ratio", "点赞率"),
    ("comment_count", "评论总数"),
    ("collected_comments", "已采集评论数"),
    ("post_type", "帖子类型"),
    ("domain", "域名"),
    ("content_url", "内容链接"),
    ("permalink", "站内路径"),
    ("full_url", "完整链接"),
    ("flair", "标签"),
    ("text", "正文"),
    ("nsfw", "是否NSFW"),
    ("media_type", "媒体类型"),
    ("media_count", "媒体数量"),
    ("media_dir", "媒体目录"),
    ("scrape_sort", "抓取排序"),
    ("scrape_time", "抓取时间"),
]

MEDIA_FIELDS = [
    ("post_id", "帖子ID"),
    ("index", "序号"),
    ("kind", "类型"),
    ("source", "来源"),
    ("url", "原始链接"),
    ("filename", "本地文件名"),
    ("path", "本地路径"),
    ("size_kb", "大小(KB)"),
    ("status", "状态"),
    ("note", "说明"),
]

MEDIA_WIDTHS = {
    "url": 60, "path": 60, "filename": 30, "source": 24, "note": 40, "status": 10,
}

COMMENT_FIELDS = [
    ("post_id", "帖子ID"),
    ("comment_id", "评论ID"),
    ("parent_id", "父评论ID"),
    ("root_comment_id", "根评论ID"),
    ("depth", "层级"),
    ("tree_no", "楼层编号"),
    ("tree_prefix", "树形结构"),
    ("position", "同级排序"),
    ("parent_positions", "父级路径"),
    ("author", "作者"),
    ("created", "发布时间"),
    ("score", "点赞数"),
    ("award_count", "奖励数"),
    ("reply_count", "回复数"),
    ("permalink", "站内路径"),
    ("full_url", "完整链接"),
    ("content_type", "内容类型"),
    ("is_op", "是否楼主"),
    ("collapsed", "是否折叠"),
    ("is_deleted", "是否已删除"),
    ("media_urls", "媒体链接"),
    ("body", "评论内容"),
]

SEARCH_FIELDS = [
    ("search_keyword", "搜索关键词"),
    ("search_subreddit", "搜索子版块"),
    ("search_sort", "搜索排序"),
    ("search_time", "时间范围"),
    ("result_rank", "结果排名"),
    ("post_id", "帖子ID"),
    ("bare_post_id", "帖子短ID"),
    ("subreddit", "子版块"),
    ("subreddit_id", "子版块ID"),
    ("title", "标题"),
    ("author", "作者"),
    ("author_fullname", "作者ID"),
    ("created_utc", "发布时间戳"),
    ("score", "点赞数"),
    ("upvote_ratio", "点赞率"),
    ("comment_count", "评论数"),
    ("flair", "标签"),
    ("domain", "域名"),
    ("is_self", "是否文字帖"),
    ("post_hint", "内容类型"),
    ("nsfw", "是否NSFW"),
    ("spoiler", "是否剧透"),
    ("locked", "是否锁定"),
    ("stickied", "是否置顶"),
    ("permalink", "站内路径"),
    ("full_url", "完整链接"),
    ("outbound_url", "外链地址"),
    ("thumbnail", "缩略图"),
    ("selftext", "正文"),
    ("distinguished", "官方标识"),
    ("search_collected_at", "采集时间"),
]


# Excel 媒体展示列（内部 key 保持英文；仅表头为中文）
def _ensure_media_field(fields, after_key, field):
    if any(k == field[0] for k, _ in fields):
        return
    for i, (k, _t) in enumerate(fields):
        if k == after_key:
            fields.insert(i + 1, field)
            return
    fields.append(field)

_ensure_media_field(POST_FIELDS, "media_dir", ("post_image", "帖子图片"))
_ensure_media_field(POST_FIELDS, "post_image", ("video_link", "视频链接"))
_ensure_media_field(MEDIA_FIELDS, "post_id", ("media_scope", "媒体归属"))
_ensure_media_field(MEDIA_FIELDS, "media_scope", ("comment_id", "评论ID"))
_ensure_media_field(MEDIA_FIELDS, "kind", ("preview", "图片预览"))

# 兼容旧引用：内部字段键列表
POST_COLUMNS = [key for key, _ in POST_FIELDS]
COMMENT_COLUMNS = [key for key, _ in COMMENT_FIELDS]
SEARCH_POST_COLUMNS = [key for key, _ in SEARCH_FIELDS]

POST_WIDTHS = {
    "title": 55, "text": 90, "content_url": 55, "full_url": 55,
    "permalink": 46, "created": 24, "scrape_time": 24, "author": 18,
}
COMMENT_WIDTHS = {
    "tree_prefix": 26, "body": 90, "full_url": 50, "permalink": 46,
    "parent_positions": 14, "created": 24, "author": 20, "tree_no": 10, "media_urls": 40,
    "depth": 6, "position": 9, "score": 9, "award_count": 9, "reply_count": 9,
}
SEARCH_WIDTHS = {
    "search_keyword": 24, "title": 60, "selftext": 90, "full_url": 60,
    "outbound_url": 60, "thumbnail": 50, "created_utc": 25, "search_collected_at": 24,
}

# 评论树绘制字符
_BRANCH = "├─ "
_LAST = "└─ "
_PASS = "│  "
_BLANK = "   "
_MAX_OUTLINE_LEVEL = 7
_MAX_INDENT = 15


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    return value


def _keys(fields: Iterable[tuple[str, str]]) -> list[str]:
    return [key for key, _ in fields]


def _decorate(ws, fields: list[tuple[str, str]], widths: dict[str, float] | None = None):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for idx, (key, title) in enumerate(fields, 1):
        width = (widths or {}).get(key, min(28, max(10, len(title) * 2 + 4)))
        ws.column_dimensions[get_column_letter(idx)].width = width


def _root_map(comments: list[dict[str, Any]]) -> tuple[dict[str, str], Counter]:
    by_id = {str(x.get("comment_id") or ""): x for x in comments}
    replies = Counter(str(x.get("parent_id") or "") for x in comments if x.get("parent_id"))
    roots: dict[str, str] = {}
    for cid, row in by_id.items():
        current = row
        seen: set[str] = set()
        root = cid
        while True:
            parent = str(current.get("parent_id") or "")
            if not parent or parent.startswith("t3_") or parent in seen or parent not in by_id:
                break
            seen.add(parent)
            root = parent
            current = by_id[parent]
        roots[cid] = root
    return roots, replies


def enrich_comments(comments: list[dict[str, Any]]) -> list[dict[str, Any]]:
    roots, replies = _root_map(comments)
    out = []
    for row in comments:
        r = dict(row)
        cid = str(r.get("comment_id") or "")
        permalink = str(r.get("permalink") or "")
        r["root_comment_id"] = roots.get(cid, cid)
        r["reply_count"] = int(replies.get(cid, 0))
        r["full_url"] = ("https://www.reddit.com" + permalink) if permalink.startswith("/") else permalink
        out.append(r)
    return out


def _sibling_key(row: dict[str, Any]):
    """同一父级下的排序：Reddit 展示顺序（position）优先，其次热度、时间。"""
    try:
        position = int(row.get("position") or 0)
    except (TypeError, ValueError):
        position = 0
    try:
        score = int(row.get("score") or 0)
    except (TypeError, ValueError):
        score = 0
    return (position, -score, str(row.get("created") or ""), str(row.get("comment_id") or ""))


def build_comment_tree(comments: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """把扁平评论列表还原成 Reddit 的评论树（深度优先，父评论紧跟其回复）。

    额外补充字段：
      depth        真实层级（0 = 顶层评论）
      tree_no      楼层编号，例如 3 / 3.2 / 3.2.1
      tree_prefix  树形引导线，例如 "│  └─ "
      root_comment_id / reply_count 由真实树结构重算
    """
    by_id: dict[str, dict[str, Any]] = {}
    for row in enrich_comments(comments):
        cid = str(row.get("comment_id") or "")
        if cid and cid not in by_id:
            by_id[cid] = row
    rows = list(by_id.values())

    children: dict[str, list[dict[str, Any]]] = {}
    roots: list[dict[str, Any]] = []
    for row in rows:
        parent_id = str(row.get("parent_id") or "")
        if parent_id.startswith("t3_"):  # 顶层评论的父级是帖子
            parent_id = ""
        parent = by_id.get(parent_id) if parent_id else None
        if parent is None or parent is row or parent_id == str(row.get("comment_id") or ""):
            roots.append(row)
        else:
            children.setdefault(parent_id, []).append(row)

    for siblings in children.values():
        siblings.sort(key=_sibling_key)

    # 兜底：父子关系成环时根集合可能为空，把环内节点降级为顶层，保证一条评论都不丢
    reachable: set[str] = set()
    pending = [str(r.get("comment_id") or "") for r in roots]
    while pending:
        cid = pending.pop()
        if cid in reachable:
            continue
        reachable.add(cid)
        pending.extend(str(k.get("comment_id") or "") for k in children.get(cid, []))
    for row in rows:
        cid = str(row.get("comment_id") or "")
        if cid and cid not in reachable:
            reachable.add(cid)
            roots.append(row)

    roots.sort(key=_sibling_key)

    ordered: list[dict[str, Any]] = []
    emitted: set[str] = set()
    stack: list[tuple[dict[str, Any], str, str, int, bool, int]] = []

    def push_siblings(nodes: list[dict[str, Any]], prefix: str, parent_no: str, depth: int):
        total = len(nodes)
        for i in range(total - 1, -1, -1):  # 逆序入栈，出栈时保持正序
            stack.append((nodes[i], prefix, parent_no, depth, i == total - 1, i + 1))

    push_siblings(roots, "", "", 0)
    while stack:
        node, prefix, parent_no, depth, is_last, index = stack.pop()
        cid = str(node.get("comment_id") or "")
        if cid and cid in emitted:  # 脏数据成环时防止无限递归
            continue
        if cid:
            emitted.add(cid)
        node["depth"] = depth
        node["tree_no"] = f"{parent_no}{index}"
        node["tree_prefix"] = prefix + (_LAST if is_last else _BRANCH)
        kids = [k for k in (children.get(cid) or []) if str(k.get("comment_id") or "") not in emitted]
        node["reply_count"] = len(kids)
        ordered.append(node)
        if kids:
            push_siblings(kids, prefix + (_BLANK if is_last else _PASS), f"{node['tree_no']}.", depth + 1)
    return ordered


def build_comment_forest(comments: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """多帖合并导出：按帖子分组后各自还原评论树。"""
    groups: dict[str, list[dict[str, Any]]] = {}
    order: list[str] = []
    for row in comments:
        pid = str(row.get("post_id") or "")
        if pid not in groups:
            groups[pid] = []
            order.append(pid)
        groups[pid].append(row)
    out: list[dict[str, Any]] = []
    for pid in order:
        out.extend(build_comment_tree(groups[pid]))
    return out


def _write_sheet(ws, fields: list[tuple[str, str]], rows: list[dict[str, Any]]):
    ws.append([title for _, title in fields])
    keys = _keys(fields)
    for row in rows:
        ws.append([_cell_value(row.get(key, "")) for key in keys])


def _apply_tree_style(ws, rows: list[dict[str, Any]]):
    """评论树可视化：缩进 + Excel 分级（可折叠/展开，和 Reddit 一样）。"""
    keys = _keys(COMMENT_FIELDS)
    col_index = {key: idx + 1 for idx, key in enumerate(keys)}
    body_col = col_index["body"]
    tree_col = col_index["tree_prefix"]
    author_col = col_index["author"]

    for offset, row in enumerate(rows):
        excel_row = offset + 2
        try:
            depth = int(row.get("depth") or 0)
        except (TypeError, ValueError):
            depth = 0
        indent = min(depth, _MAX_INDENT)
        if depth > 0:
            ws.row_dimensions[excel_row].outline_level = min(depth, _MAX_OUTLINE_LEVEL)
        ws.cell(excel_row, tree_col).alignment = Alignment(vertical="top", wrap_text=False)
        ws.cell(excel_row, author_col).alignment = Alignment(vertical="top", wrap_text=True, indent=indent)
        ws.cell(excel_row, body_col).alignment = Alignment(vertical="top", wrap_text=True, indent=indent)

    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False)


def _write_comment_sheet(ws, comments: list[dict[str, Any]], multi_post: bool = False):
    rows = build_comment_forest(comments) if multi_post else build_comment_tree(comments)
    widths = dict(COMMENT_WIDTHS)
    max_depth = max((int(r.get("depth") or 0) for r in rows), default=0)
    widths["tree_prefix"] = min(60, max(12, max_depth * 3 + 4))
    _write_sheet(ws, COMMENT_FIELDS, rows)
    _decorate(ws, COMMENT_FIELDS, widths)
    _apply_tree_style(ws, rows)
    return rows


def _write_media_sheet(ws, media_rows: list[dict[str, Any]]):
    _write_sheet(ws, MEDIA_FIELDS, media_rows)
    _decorate(ws, MEDIA_FIELDS, MEDIA_WIDTHS)
    embed_media_sheet(ws, media_rows, MEDIA_FIELDS)


def write_post_workbook(
    path: Path,
    post: dict[str, Any],
    comments: list[dict[str, Any]],
    media_rows: list[dict[str, Any]] | None = None,
):
    _public_write_detail_workbook(
        path,
        [post],
        comments,
        media_rows or [],
        multi_post=False,
    )


def write_combined_workbook(
    path: Path,
    posts: list[dict[str, Any]],
    comments: list[dict[str, Any]],
    media_rows: list[dict[str, Any]] | None = None,
):
    _public_write_detail_workbook(
        path,
        posts,
        comments,
        media_rows or [],
        multi_post=True,
    )


def write_search_workbook(path: Path, rows: list[dict[str, Any]]):
    public_rows = _public_search_rows(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "搜索结果"
    _public_write_table(
        ws,
        _PUBLIC_SEARCH_FIELDS,
        public_rows,
        _PUBLIC_SEARCH_WIDTHS,
    )
    wb.save(path)
def write_combined_search_workbook(path: Path, rows: list[dict[str, Any]]):
    write_search_workbook(path, rows)


# ======================================================================================
# Public-facing Excel export schema
# Internal scraper/checkpoint fields remain unchanged. Only the workbook presentation is simplified.
# ======================================================================================

_PUBLIC_POST_FIELDS = [
    ("subreddit", "子版块"),
    ("title", "标题"),
    ("author", "作者"),
    ("created", "发布时间"),
    ("score", "点赞数"),
    ("upvote_ratio", "点赞率"),
    ("comment_count", "评论数"),
    ("flair", "标签"),
    ("text", "正文"),
    ("post_image", "帖子图片"),
    ("video_link", "视频链接"),
    ("full_url", "帖子链接"),
    ("content_url_public", "外部内容链接"),
]

_PUBLIC_COMMENT_FIELDS = [
    ("post_title", "帖子标题"),
    ("tree_no", "楼层"),
    ("author", "作者"),
    ("created", "发布时间"),
    ("score", "点赞数"),
    ("reply_count", "回复数"),
    ("is_op", "是否楼主"),
    ("body", "评论内容"),
    ("full_url", "评论链接"),
]

_PUBLIC_POST_MEDIA_FIELDS = [
    ("post_title", "帖子标题"),
    ("display_index", "序号"),
    ("preview", "图片预览"),
    ("media_type_public", "媒体类型"),
    ("url", "媒体链接"),
]

_PUBLIC_COMMENT_MEDIA_FIELDS = [
    ("post_title", "帖子标题"),
    ("comment_author", "评论作者"),
    ("comment_body", "评论内容"),
    ("preview", "图片预览"),
    ("media_type_public", "媒体类型"),
    ("url", "媒体链接"),
    ("comment_url", "评论链接"),
]

_PUBLIC_SEARCH_FIELDS = [
    ("search_keyword", "搜索关键词"),
    ("search_scope_public", "搜索范围"),
    ("search_sort", "排序方式"),
    ("search_time", "时间范围"),
    ("result_rank", "排名"),
    ("subreddit", "子版块"),
    ("title", "标题"),
    ("author", "作者"),
    ("created_utc", "发布时间"),
    ("score", "点赞数"),
    ("upvote_ratio", "点赞率"),
    ("comment_count", "评论数"),
    ("flair", "标签"),
    ("selftext", "正文"),
    ("full_url", "帖子链接"),
    ("outbound_url_public", "外部链接"),
]

_PUBLIC_POST_WIDTHS = {
    "subreddit": 20, "title": 56, "author": 18, "created": 22,
    "score": 10, "upvote_ratio": 10, "comment_count": 10, "flair": 18,
    "text": 72, "post_image": 28, "video_link": 46,
    "full_url": 48, "content_url_public": 48,
}

_PUBLIC_COMMENT_WIDTHS = {
    "post_title": 42, "tree_no": 10, "author": 18, "created": 22,
    "score": 10, "reply_count": 10, "is_op": 10, "body": 80, "full_url": 48,
}

_PUBLIC_POST_MEDIA_WIDTHS = {
    "post_title": 42, "display_index": 8, "preview": 28,
    "media_type_public": 12, "url": 56,
}

_PUBLIC_COMMENT_MEDIA_WIDTHS = {
    "post_title": 40, "comment_author": 18, "comment_body": 60,
    "preview": 28, "media_type_public": 12, "url": 54, "comment_url": 48,
}

_PUBLIC_SEARCH_WIDTHS = {
    "search_keyword": 24, "search_scope_public": 18, "search_sort": 12,
    "search_time": 12, "result_rank": 8, "subreddit": 20, "title": 56,
    "author": 18, "created_utc": 22, "score": 10, "upvote_ratio": 10,
    "comment_count": 10, "flair": 18, "selftext": 72,
    "full_url": 54, "outbound_url_public": 54,
}


def _public_bare_post_id(value: Any) -> str:
    s = str(value or "")
    return s[3:] if s.startswith("t3_") else s


def _public_media_scope(row: dict[str, Any]) -> str:
    scope = str(row.get("media_scope") or "").strip().lower()
    if scope in {"post", "comment"}:
        return scope
    return "comment" if str(row.get("comment_id") or "").strip() else "post"


def _public_post_lookup(posts: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for row in posts:
        pid = _public_bare_post_id(row.get("post_id"))
        if pid:
            out[pid] = row
    return out


def _public_comment_lookup(comments: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for row in comments:
        cid = str(row.get("comment_id") or "")
        if cid:
            out[cid] = row
    return out


def _public_split_media(media_rows: list[dict[str, Any]]):
    post_media: list[dict[str, Any]] = []
    comment_media: list[dict[str, Any]] = []
    for row in media_rows or []:
        if _public_media_scope(row) == "comment":
            comment_media.append(row)
        else:
            post_media.append(row)
    return post_media, comment_media


def _public_external_content_url(post: dict[str, Any]) -> str:
    full_url = str(post.get("full_url") or "")
    content_url = str(post.get("content_url") or "")
    if not content_url.startswith("http"):
        return ""
    # Reddit-hosted image/video belongs in image/video columns, not as a duplicate "external link".
    low = content_url.lower()
    if any(host in low for host in ("reddit.com/", "redd.it/", "preview.redd.it/", "v.redd.it/")):
        return ""
    if content_url == full_url:
        return ""
    return content_url


def _public_outbound_url(search_row: dict[str, Any]) -> str:
    value = str(search_row.get("outbound_url") or "")
    full_url = str(search_row.get("full_url") or "")
    if not value.startswith("http") or value == full_url:
        return ""
    return value


def _public_post_rows(
    posts: list[dict[str, Any]],
    media_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    post_media, _ = _public_split_media(media_rows)
    grouped: dict[str, list[dict[str, Any]]] = {}
    for m in post_media:
        pid = _public_bare_post_id(m.get("post_id"))
        grouped.setdefault(pid, []).append(m)

    out: list[dict[str, Any]] = []
    for src in posts:
        row = dict(src)
        pid = _public_bare_post_id(row.get("post_id"))
        related = grouped.get(pid, [])

        image_path = ""
        video_link = ""
        for m in related:
            kind = str(m.get("kind") or "").lower()
            status = str(m.get("status") or "")
            if kind == "image" and not image_path and status in {"成功", "已存在"}:
                p = str(m.get("path") or "")
                if p and Path(p).is_file():
                    image_path = p
            if kind == "video" and not video_link:
                u = str(m.get("url") or "")
                if u.startswith("http"):
                    video_link = u

        row["post_image"] = ""
        row["__post_image_path"] = image_path
        row["video_link"] = video_link
        row["content_url_public"] = _public_external_content_url(row)
        out.append(row)

    return out


def _public_comment_rows(
    comments: list[dict[str, Any]],
    posts: list[dict[str, Any]],
    multi_post: bool,
) -> list[dict[str, Any]]:
    post_map = _public_post_lookup(posts)
    ordered = build_comment_forest(comments) if multi_post else build_comment_tree(comments)

    out: list[dict[str, Any]] = []
    for src in ordered:
        row = dict(src)
        pid = _public_bare_post_id(row.get("post_id"))
        post = post_map.get(pid) or {}
        row["post_title"] = str(post.get("title") or "")
        out.append(row)
    return out


def _public_media_type(kind: Any) -> str:
    k = str(kind or "").lower()
    if k == "image":
        return "图片"
    if k == "video":
        return "视频"
    return "其他"


def _public_post_media_rows(
    media_rows: list[dict[str, Any]],
    posts: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    post_map = _public_post_lookup(posts)
    post_media, _ = _public_split_media(media_rows)

    out: list[dict[str, Any]] = []
    counters: dict[str, int] = {}
    for src in post_media:
        row = dict(src)
        pid = _public_bare_post_id(row.get("post_id"))
        counters[pid] = counters.get(pid, 0) + 1
        post = post_map.get(pid) or {}
        row["post_title"] = str(post.get("title") or "")
        row["display_index"] = counters[pid]
        row["preview"] = ""
        row["media_type_public"] = _public_media_type(row.get("kind"))
        out.append(row)
    return out


def _public_comment_media_rows(
    media_rows: list[dict[str, Any]],
    posts: list[dict[str, Any]],
    comments: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    post_map = _public_post_lookup(posts)
    comment_map = _public_comment_lookup(comments)
    _, comment_media = _public_split_media(media_rows)

    out: list[dict[str, Any]] = []
    for src in comment_media:
        row = dict(src)
        pid = _public_bare_post_id(row.get("post_id"))
        cid = str(row.get("comment_id") or "")
        post = post_map.get(pid) or {}
        comment = comment_map.get(cid) or {}

        row["post_title"] = str(post.get("title") or "")
        row["comment_author"] = str(comment.get("author") or "")
        row["comment_body"] = str(comment.get("body") or "")
        row["comment_url"] = str(comment.get("full_url") or "")
        if not row["comment_url"]:
            permalink = str(comment.get("permalink") or "")
            if permalink.startswith("/"):
                row["comment_url"] = "https://www.reddit.com" + permalink
            else:
                row["comment_url"] = permalink

        row["preview"] = ""
        row["media_type_public"] = _public_media_type(row.get("kind"))
        out.append(row)
    return out


def _public_search_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for src in rows:
        row = dict(src)
        scope = str(row.get("search_subreddit") or "").strip()
        row["search_scope_public"] = ("r/" + scope.lstrip("r/")) if scope else "全站"
        row["outbound_url_public"] = _public_outbound_url(row)
        out.append(row)
    return out


def _public_write_table(
    ws,
    fields: list[tuple[str, str]],
    rows: list[dict[str, Any]],
    widths: dict[str, float],
):
    _write_sheet(ws, fields, rows)
    _decorate(ws, fields, widths)

    key_to_col = {key: idx + 1 for idx, (key, _title) in enumerate(fields)}

    # Readable numbers / percentages.
    if "upvote_ratio" in key_to_col:
        col = key_to_col["upvote_ratio"]
        for r in range(2, ws.max_row + 1):
            value = ws.cell(r, col).value
            if isinstance(value, (int, float)):
                ws.cell(r, col).number_format = "0.0%"

    # Clickable URLs without exposing local filesystem paths.
    for key in ("video_link", "full_url", "content_url_public", "url", "comment_url", "outbound_url_public"):
        col = key_to_col.get(key)
        if not col:
            continue
        for r in range(2, ws.max_row + 1):
            value = ws.cell(r, col).value
            if isinstance(value, str) and value.startswith("http"):
                ws.cell(r, col).hyperlink = value
                ws.cell(r, col).style = "Hyperlink"


def _public_add_image(ws, cell_coordinate: str, path: str, max_w: int = 180, max_h: int = 120):
    try:
        from io import BytesIO
        from openpyxl.drawing.image import Image as XLImage
        from PIL import Image as PILImage

        with PILImage.open(path) as im0:
            im = im0.copy()

        if im.mode not in {"RGB", "RGBA"}:
            im = im.convert("RGBA")

        im.thumbnail((max_w, max_h))
        bio = BytesIO()
        im.save(bio, format="PNG", optimize=True)
        bio.seek(0)

        pic = XLImage(bio)
        pic.width, pic.height = im.size
        ws.add_image(pic, cell_coordinate)
        return True
    except Exception:
        return False


def _public_embed_post_images(ws, rows: list[dict[str, Any]]):
    fields = _PUBLIC_POST_FIELDS
    image_col = next((i for i, (k, _t) in enumerate(fields, 1) if k == "post_image"), 0)
    if not image_col:
        return

    for idx, row in enumerate(rows, 2):
        p = str(row.get("__post_image_path") or "")
        if p and Path(p).is_file():
            if _public_add_image(ws, ws.cell(idx, image_col).coordinate, p):
                ws.row_dimensions[idx].height = max(ws.row_dimensions[idx].height or 15, 90)


def _public_embed_media_images(ws, rows: list[dict[str, Any]], fields: list[tuple[str, str]]):
    preview_col = next((i for i, (k, _t) in enumerate(fields, 1) if k == "preview"), 0)
    if not preview_col:
        return

    for idx, row in enumerate(rows, 2):
        if str(row.get("kind") or "").lower() != "image":
            continue
        if str(row.get("status") or "") not in {"成功", "已存在"}:
            continue

        p = str(row.get("path") or "")
        if p and Path(p).is_file():
            if _public_add_image(ws, ws.cell(idx, preview_col).coordinate, p):
                ws.row_dimensions[idx].height = max(ws.row_dimensions[idx].height or 15, 90)


def _public_style_comment_depth(ws, rows: list[dict[str, Any]]):
    body_col = next((i for i, (k, _t) in enumerate(_PUBLIC_COMMENT_FIELDS, 1) if k == "body"), 0)
    author_col = next((i for i, (k, _t) in enumerate(_PUBLIC_COMMENT_FIELDS, 1) if k == "author"), 0)
    if not body_col:
        return

    for idx, row in enumerate(rows, 2):
        try:
            depth = min(int(row.get("depth") or 0), 12)
        except Exception:
            depth = 0
        ws.cell(idx, body_col).alignment = Alignment(vertical="top", wrap_text=True, indent=depth)
        if author_col:
            ws.cell(idx, author_col).alignment = Alignment(vertical="top", wrap_text=True, indent=depth)


def _public_write_detail_workbook(
    path: Path,
    posts: list[dict[str, Any]],
    comments: list[dict[str, Any]],
    media_rows: list[dict[str, Any]],
    multi_post: bool,
):
    post_rows = _public_post_rows(posts, media_rows)
    comment_rows = _public_comment_rows(comments, posts, multi_post)
    post_media_rows = _public_post_media_rows(media_rows, posts)
    comment_media_rows = _public_comment_media_rows(media_rows, posts, comments)

    wb = Workbook()

    ws_post = wb.active
    ws_post.title = "帖子"
    _public_write_table(ws_post, _PUBLIC_POST_FIELDS, post_rows, _PUBLIC_POST_WIDTHS)
    _public_embed_post_images(ws_post, post_rows)

    ws_comments = wb.create_sheet("评论")
    _public_write_table(ws_comments, _PUBLIC_COMMENT_FIELDS, comment_rows, _PUBLIC_COMMENT_WIDTHS)
    _public_style_comment_depth(ws_comments, comment_rows)

    ws_post_media = wb.create_sheet("帖子媒体")
    _public_write_table(
        ws_post_media,
        _PUBLIC_POST_MEDIA_FIELDS,
        post_media_rows,
        _PUBLIC_POST_MEDIA_WIDTHS,
    )
    _public_embed_media_images(ws_post_media, post_media_rows, _PUBLIC_POST_MEDIA_FIELDS)

    ws_comment_media = wb.create_sheet("评论媒体")
    _public_write_table(
        ws_comment_media,
        _PUBLIC_COMMENT_MEDIA_FIELDS,
        comment_media_rows,
        _PUBLIC_COMMENT_MEDIA_WIDTHS,
    )
    _public_embed_media_images(ws_comment_media, comment_media_rows, _PUBLIC_COMMENT_MEDIA_FIELDS)

    wb.save(path)
# End of file.
