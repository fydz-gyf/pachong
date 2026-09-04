from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage


def _bare(value: Any) -> str:
    s = str(value or "")
    return s[3:] if s.startswith("t3_") else s


def media_scope(row: dict[str, Any]) -> str:
    scope = str(row.get("media_scope") or "").strip().lower()
    if scope in {"post", "comment"}:
        return scope
    if str(row.get("comment_id") or "").strip():
        return "comment"
    return "post"


def split_media_rows(media_rows: list[dict[str, Any]]):
    post_rows = []
    comment_rows = []
    for row in media_rows:
        if media_scope(row) == "comment":
            comment_rows.append(row)
        else:
            post_rows.append(row)
    return post_rows, comment_rows


def _rows_for_post(media_rows: list[dict[str, Any]], post_id: Any):
    bare = _bare(post_id)
    return [
        r for r in media_rows
        if media_scope(r) == "post" and _bare(r.get("post_id")) == bare
    ]


def prepare_posts_for_media(posts, media_rows):
    out = []
    for src in posts:
        row = dict(src)
        related = _rows_for_post(media_rows, row.get("post_id"))

        image_path = ""
        for m in related:
            path = str(m.get("path") or "").strip()
            if (
                m.get("kind") == "image"
                and m.get("status") in {"成功", "已存在"}
                and path
                and Path(path).is_file()
            ):
                image_path = path
                break

        video_link = ""
        for m in related:
            if m.get("kind") == "video":
                link = str(m.get("url") or "").strip()
                if link.startswith("http"):
                    video_link = link
                    break

        row["post_image"] = ""
        row["video_link"] = video_link
        row["__post_image_path"] = image_path
        out.append(row)
    return out


def _field_col(fields, key):
    for i, (k, _title) in enumerate(fields, 1):
        if k == key:
            return i
    return 0


def _xl_thumb(path: str, max_w: int = 180, max_h: int = 120):
    try:
        with PILImage.open(path) as im0:
            im = im0.copy()
        if im.mode not in {"RGB", "RGBA"}:
            im = im.convert("RGBA")
        im.thumbnail((max_w, max_h))
        bio = BytesIO()
        im.save(bio, format="PNG", optimize=True)
        bio.seek(0)
        pic = XLImage(bio)
        pic.width, pic.height = im.size
        return pic
    except Exception:
        return None


def embed_post_sheet(ws, posts, fields):
    img_col = _field_col(fields, "post_image")
    video_col = _field_col(fields, "video_link")
    if img_col:
        ws.column_dimensions[ws.cell(1, img_col).column_letter].width = 28
    if video_col:
        ws.column_dimensions[ws.cell(1, video_col).column_letter].width = 55

    for idx, row in enumerate(posts, 2):
        if img_col:
            path = str(row.get("__post_image_path") or "")
            if path and Path(path).is_file():
                pic = _xl_thumb(path)
                if pic is not None:
                    ws.add_image(pic, ws.cell(idx, img_col).coordinate)
                    ws.row_dimensions[idx].height = max(
                        ws.row_dimensions[idx].height or 15, 90
                    )
        if video_col:
            link = str(row.get("video_link") or "")
            if link.startswith("http"):
                cell = ws.cell(idx, video_col)
                cell.hyperlink = link
                cell.style = "Hyperlink"


def embed_media_sheet(ws, media_rows, fields):
    preview_col = _field_col(fields, "preview")
    url_col = _field_col(fields, "url")
    if preview_col:
        ws.column_dimensions[ws.cell(1, preview_col).column_letter].width = 28

    for idx, row in enumerate(media_rows, 2):
        if (
            row.get("kind") == "image"
            and row.get("status") in {"成功", "已存在"}
        ):
            path = str(row.get("path") or "")
            if preview_col and path and Path(path).is_file():
                pic = _xl_thumb(path)
                if pic is not None:
                    ws.add_image(pic, ws.cell(idx, preview_col).coordinate)
                    ws.row_dimensions[idx].height = max(
                        ws.row_dimensions[idx].height or 15, 90
                    )
        if url_col:
            link = str(row.get("url") or "")
            if link.startswith("http"):
                cell = ws.cell(idx, url_col)
                cell.hyperlink = link
                cell.style = "Hyperlink"
