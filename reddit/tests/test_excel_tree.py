from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl

from reddit_scraper.storage.excel import (
    _PUBLIC_COMMENT_FIELDS,
    _PUBLIC_COMMENT_MEDIA_FIELDS,
    _PUBLIC_POST_FIELDS,
    _PUBLIC_POST_MEDIA_FIELDS,
    build_comment_forest,
    build_comment_tree,
    write_combined_workbook,
    write_post_workbook,
)


def _c(cid, parent="", depth=0, position=0, score=0, post="t3_p1", author="u"):
    return {
        "post_id": post, "comment_id": cid, "parent_id": parent, "depth": depth,
        "position": position, "parent_positions": "[]", "author": author + cid[-1],
        "created": "2026-01-01T00:00:00+0000", "score": score, "award_count": 0,
        "permalink": f"/r/x/comments/p1/comment/{cid}/", "content_type": "text",
        "is_op": False, "collapsed": False, "is_deleted": False, "body": f"body-{cid}",
    }


# root a (score 10) -> b (score 5) -> c ; root d (score 1)
COMMENTS = [
    _c("t1_a", parent="", position=0, score=10),
    _c("t1_b", parent="t1_a", depth=1, position=0, score=5),
    _c("t1_c", parent="t1_b", depth=2, position=0, score=2),
    _c("t1_d", parent="t3_p1", position=1, score=1),  # 顶层评论父级可能是帖子
    _c("t1_orphan", parent="t1_missing", depth=4, position=3, score=0),  # 父评论未采集
]

tree = build_comment_tree(COMMENTS)
order = [r["comment_id"] for r in tree]
assert order == ["t1_a", "t1_b", "t1_c", "t1_d", "t1_orphan"], order

a, b, c, d, orphan = tree
assert (a["depth"], b["depth"], c["depth"]) == (0, 1, 2)
assert (a["tree_no"], b["tree_no"], c["tree_no"]) == ("1", "1.1", "1.1.1")
assert a["tree_prefix"] == "├─ ", a["tree_prefix"]  # 顶层非最后一条
assert b["tree_prefix"] == "│  └─ ", b["tree_prefix"]  # a 的唯一回复
assert c["tree_prefix"] == "│     └─ ", c["tree_prefix"]  # 再深一级，父级已收尾
assert (a["reply_count"], b["reply_count"], c["reply_count"]) == (1, 1, 0)
assert {a["root_comment_id"], b["root_comment_id"], c["root_comment_id"]} == {"t1_a"}
# 父评论未采集时降级为顶层，不能丢数据
assert orphan["depth"] == 0 and orphan["root_comment_id"] == "t1_orphan"
# 全量覆盖，无遗漏无重复
assert len(tree) == len(COMMENTS)

# 多帖合并：每个帖子各自成树
forest = build_comment_forest(COMMENTS + [_c("t1_z", post="t3_p2")])
assert [r["comment_id"] for r in forest] == ["t1_a", "t1_b", "t1_c", "t1_d", "t1_orphan", "t1_z"]

# 空输入
assert build_comment_tree([]) == []

with tempfile.TemporaryDirectory() as tmp:
    out = Path(tmp) / "post.xlsx"
    write_post_workbook(out, {"post_id": "t3_p1", "title": "T", "nsfw": True}, COMMENTS)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["帖子", "评论", "帖子媒体", "评论媒体"], wb.sheetnames
    post_headers = [cell.value for cell in wb["帖子"][1]]
    assert post_headers == [label for _, label in _PUBLIC_POST_FIELDS], post_headers
    headers = [cell.value for cell in wb["评论"][1]]
    assert headers == [label for _, label in _PUBLIC_COMMENT_FIELDS], headers
    assert headers[0] == "帖子标题" and "评论内容" in headers and "评论链接" in headers
    body_col = headers.index("评论内容") + 1
    assert wb["评论"].cell(3, body_col).alignment.indent == 1  # 回复缩进一级
    assert [cell.value for cell in wb["帖子媒体"][1]] == [label for _, label in _PUBLIC_POST_MEDIA_FIELDS]
    assert [cell.value for cell in wb["评论媒体"][1]] == [label for _, label in _PUBLIC_COMMENT_MEDIA_FIELDS]

    combined = Path(tmp) / "combined.xlsx"
    write_combined_workbook(combined, [{"post_id": "t3_p1"}], COMMENTS)
    assert openpyxl.load_workbook(combined).sheetnames == ["帖子", "评论", "帖子媒体", "评论媒体"]

print("TESTS: OK")
