"""Excel 导出。"""

from __future__ import annotations

import logging
import os
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import Settings
from .images import download_keyword_images
from .utils import normalize_url, safe_filename, sanitize_sheet_name, to_number

EXCEL_HEADERS = [
    "总排名",
    "页码",
    "页内排名",
    "接口位置",
    "主图",
    "商品ID",
    "商品名称",
    "前台价格",
    "原价",
    "价格说明",
    "销量/热度",
    "店铺名称",
    "卖家昵称",
    "发货地区",
    "类目ID",
    "是否广告",
    "是否有视频",
    "商品链接",
    "视频链接",
    "主图URL",
    "视频封面URL",
]

COLUMN_WIDTHS = {
    "A": 9,
    "B": 7,
    "C": 9,
    "D": 9,
    "E": 18,
    "F": 18,
    "G": 52,
    "H": 13,
    "I": 13,
    "J": 12,
    "K": 16,
    "L": 24,
    "M": 22,
    "N": 16,
    "O": 14,
    "P": 11,
    "Q": 12,
    "R": 13,
    "S": 13,
    "T": 14,
    "U": 16,
}

# 主图列的行高（磅）。px -> pt 约 0.75，再加一点余量。
IMAGE_ROW_HEIGHT_RATIO = 0.75
TEXT_ROW_HEIGHT = 42


def set_link(cell, url: str, text: str) -> None:
    url = normalize_url(url)
    if not url:
        cell.value = ""
        return
    cell.value = text
    cell.hyperlink = url
    cell.font = Font(color="0563C1", underline="single")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def image_row_height(settings: Settings) -> int:
    return int(settings.image_max_size * IMAGE_ROW_HEIGHT_RATIO) + 2


def style_sheet(ws, row_count: int, settings: Settings) -> None:
    ws.freeze_panes = "F2"  # 冻结表头 + A:E
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(EXCEL_HEADERS))}{max(1, row_count + 1)}"
    )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    header_align = Alignment(horizontal="center", vertical="center")

    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    alt_fill = PatternFill("solid", fgColor="F6F9FC")
    video_fill = PatternFill("solid", fgColor="E2F0D9")
    ad_fill = PatternFill("solid", fgColor="FFF2CC")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border
    ws.row_dimensions[1].height = 28

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width

    row_height = image_row_height(settings) if settings.embed_images else TEXT_ROW_HEIGHT

    for row in range(2, row_count + 2):
        ws.row_dimensions[row].height = row_height

        for col in range(1, len(EXCEL_HEADERS) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if row % 2 == 0:
                cell.fill = alt_fill

        for col in ["A", "B", "C", "D", "P", "Q"]:
            ws[f"{col}{row}"].alignment = Alignment(
                horizontal="center", vertical="center"
            )

        ws[f"F{row}"].number_format = "@"  # 商品ID文本
        ws[f"O{row}"].number_format = "@"  # 类目ID文本
        ws[f"H{row}"].number_format = "¥#,##0.00"
        ws[f"I{row}"].number_format = "¥#,##0.00"
        ws[f"G{row}"].alignment = Alignment(vertical="center", wrap_text=True)
        ws[f"L{row}"].alignment = Alignment(vertical="center", wrap_text=True)
        ws[f"M{row}"].alignment = Alignment(vertical="center", wrap_text=True)

        if str(ws[f"P{row}"].value).upper() == "Y":
            ws[f"P{row}"].fill = ad_fill
        if str(ws[f"Q{row}"].value).upper() == "Y":
            ws[f"Q{row}"].fill = video_fill


def build_excel_filename(keywords, settings: Settings) -> str:
    """根据本次输入关键词生成 Excel 文件名。

    单关键词：鼠标.xlsx
    多关键词：鼠标、键盘、耳机.xlsx
    开启 excel_append_timestamp 时追加时间戳，避免重跑覆盖上次结果。
    """
    cleaned = []
    for keyword in keywords:
        name = str(keyword).strip()
        if not name:
            continue
        # 保留关键词中的空格，只替换 Windows 文件名非法字符。
        name = re.sub(r'[\\/:*?"<>|]+', "_", name)
        name = name.rstrip(". ")
        if name:
            cleaned.append(name)

    if not cleaned:
        base = "淘宝搜索结果"
    elif len(cleaned) == 1:
        base = cleaned[0]
    else:
        base = "、".join(cleaned)

    # 给 .xlsx、时间戳和路径留余量，避免 Windows 文件名过长。
    max_base = 150
    if settings.excel_append_timestamp:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        base = f"{base[:max_base - len(stamp) - 1].rstrip('. 、_')}_{stamp}"
    elif len(base) > max_base:
        base = base[:max_base].rstrip(". 、_")

    return f"{base}.xlsx"


def _fill_sheet(
    ws,
    keyword: str,
    products: list[dict],
    environment: dict,
    settings: Settings,
) -> Path | None:
    """写入一个工作表，返回需要清理的图片缓存目录。"""
    ws.append(EXCEL_HEADERS)

    image_map = download_keyword_images(
        keyword=keyword,
        products=products,
        user_agent=environment["userAgent"],
        settings=settings,
    )
    cache_dir = (
        settings.image_cache_dir / safe_filename(keyword)
        if settings.embed_images
        else None
    )

    for row_idx, p in enumerate(products, start=2):
        values = [
            p.get("global_rank", row_idx - 1),
            p.get("page", ""),
            p.get("page_rank", ""),
            p.get("raw_position", ""),
            "",  # E 主图
            str(p.get("item_id", "")),
            p.get("title", ""),
            to_number(p.get("price", "")),
            to_number(p.get("original_price", "")),
            p.get("price_desc", ""),
            p.get("sales", ""),
            p.get("shop_name", ""),
            p.get("seller_nick", ""),
            p.get("location", ""),
            str(p.get("category", "")),
            "Y" if p.get("is_ad") else "N",
            ("Y" if p.get("has_video") else "N") if settings.export_video else "",
            "",
            "",
            "",
            "",
        ]
        ws.append(values)

        ws[f"F{row_idx}"].number_format = "@"
        ws[f"O{row_idx}"].number_format = "@"

        set_link(ws[f"R{row_idx}"], p.get("product_url", ""), "查看商品")
        if settings.export_video:
            set_link(ws[f"S{row_idx}"], p.get("video_url", ""), "查看视频")
        else:
            ws[f"S{row_idx}"].value = ""
        set_link(ws[f"T{row_idx}"], p.get("image_url", ""), "查看原图")
        if settings.export_video:
            set_link(ws[f"U{row_idx}"], p.get("video_cover", ""), "视频封面")
        else:
            ws[f"U{row_idx}"].value = ""

        image_path = image_map.get(str(p.get("item_id", "")), "")
        if image_path and os.path.exists(image_path):
            try:
                img = ExcelImage(image_path)
                img.width = settings.image_max_size
                img.height = settings.image_max_size
                img.anchor = f"E{row_idx}"
                ws.add_image(img)
            except Exception as e:
                logging.warning(f"Excel嵌图失败 item_id={p.get('item_id')}: {e}")

    style_sheet(ws, len(products), settings)
    if not settings.embed_images:
        # 主图嵌入关闭时隐藏空白的"主图"列；主图URL仍保留用于后续处理。
        ws.column_dimensions["E"].hidden = True
    if not settings.export_video:
        # 不导出视频时隐藏：是否有视频 / 视频链接 / 视频封面URL
        ws.column_dimensions["Q"].hidden = True
        ws.column_dimensions["S"].hidden = True
        ws.column_dimensions["U"].hidden = True

    return cache_dir


def export_excel(
    all_results: dict[str, list[dict]],
    environment: dict,
    settings: Settings,
) -> Path:
    output_path = settings.output_dir / build_excel_filename(
        all_results.keys(), settings
    )

    wb = Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()
    cache_dirs: list[Path] = []

    for keyword, products in all_results.items():
        sheet_name = sanitize_sheet_name(keyword, used_names)
        ws = wb.create_sheet(sheet_name)
        try:
            cache_dir = _fill_sheet(ws, keyword, products, environment, settings)
        except Exception as e:
            # 单个关键词导出失败不该让整份 Excel 丢失
            logging.exception(f"[{keyword}] Sheet 生成失败，该表仅保留表头: {e}")
            cache_dir = None

        if cache_dir is not None:
            cache_dirs.append(cache_dir)
        logging.info(f"Excel Sheet [{sheet_name}] 完成：{len(products)} 条")

    if not wb.sheetnames:
        ws = wb.create_sheet("无数据")
        ws.append(EXCEL_HEADERS)
        style_sheet(ws, 0, settings)

    wb.save(output_path)
    logging.info(f"Excel 已保存: {output_path}")

    if settings.embed_images and not settings.keep_image_cache:
        for d in cache_dirs:
            try:
                shutil.rmtree(d, ignore_errors=True)
            except Exception:
                pass

    return output_path
