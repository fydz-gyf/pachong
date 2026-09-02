from __future__ import annotations

import logging
import os
import re
import shutil

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import Settings
from ..media.images import ImageDownloader
from ..utils import default_user_agent, safe_filename, sanitize_sheet_name, to_number

HEADERS = ["总排名", "页码", "页内排名", "页面位置", "主图", "Walmart商品ID", "商品名称", "现价(USD)", "原价(USD)", "评分", "评论数", "品牌", "卖家", "库存状态", "配送/履约信息", "是否广告", "商品卡类型", "商品链接", "主图URL"]


class ExcelExporter:
    def __init__(self, settings: Settings, images: ImageDownloader):
        self.settings = settings
        self.images = images

    def _style(self, ws, row_count: int):
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="D9E2F3")
        ad_fill = PatternFill("solid", fgColor="FFF2CC")
        for cell in ws[1]:
            cell.fill = header_fill; cell.font = header_font; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        widths = [10, 8, 10, 10, 18, 18, 55, 12, 12, 10, 12, 18, 25, 16, 45, 10, 18, 14, 45]
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:S{max(1, row_count + 1)}"; ws.row_dimensions[1].height = 30
        for row in range(2, row_count + 2):
            ws.row_dimensions[row].height = max(32, self.settings.image_max_size * 0.78 if self.settings.embed_images else 32)
            ws[f"F{row}"].number_format = "@"; ws[f"H{row}"].number_format = '$#,##0.00'; ws[f"I{row}"].number_format = '$#,##0.00'
            ws[f"G{row}"].alignment = Alignment(vertical="center", wrap_text=True); ws[f"O{row}"].alignment = Alignment(vertical="center", wrap_text=True)
            if str(ws[f"P{row}"].value).upper() == "Y": ws[f"P{row}"].fill = ad_fill

    @staticmethod
    def _set_link(cell, url: str, text: str):
        if url:
            cell.value = text; cell.hyperlink = url; cell.style = "Hyperlink"

    @staticmethod
    def _filename(keywords):
        cleaned = [re.sub(r'[\\/:*?"<>|]+', "_", str(k).strip()).rstrip(". ") for k in keywords if str(k).strip()]
        base = cleaned[0] if len(cleaned) == 1 else "、".join(cleaned) if cleaned else "Walmart搜索结果"
        return f"{base[:150].rstrip('. 、_')}.xlsx"

    def export(self, all_results: dict[str, list[dict]], environment: dict):
        output = self.settings.paths.output / self._filename(all_results.keys())
        wb = Workbook(); wb.remove(wb.active); used = set(); cleanup = []
        for keyword, products in all_results.items():
            ws = wb.create_sheet(sanitize_sheet_name(keyword, used)); ws.append(HEADERS)
            image_map = self.images.download_keyword(keyword, products, environment.get("userAgent", default_user_agent()))
            if self.settings.embed_images: cleanup.append(self.settings.paths.images / safe_filename(keyword))
            for row_idx, p in enumerate(products, start=2):
                ws.append([p.get("global_rank", row_idx-1), p.get("page", ""), p.get("page_rank", ""), p.get("raw_position", ""), "", str(p.get("item_id", "")), p.get("title", ""), to_number(p.get("price", "")), to_number(p.get("original_price", "")), to_number(p.get("rating", "")), to_number(p.get("review_count", "")), p.get("brand", ""), p.get("seller", ""), p.get("availability", ""), p.get("fulfillment", ""), "Y" if p.get("is_sponsored") else "N", p.get("raw_type", ""), "", ""])
                self._set_link(ws[f"R{row_idx}"], p.get("product_url", ""), "查看商品"); self._set_link(ws[f"S{row_idx}"], p.get("image_url", ""), "查看原图")
                image_path = image_map.get(str(p.get("item_id", "")), "")
                if image_path and os.path.exists(image_path):
                    try:
                        img = ExcelImage(image_path); img.width = self.settings.image_max_size; img.height = self.settings.image_max_size; img.anchor = f"E{row_idx}"; ws.add_image(img)
                    except Exception as e:
                        logging.warning("Excel image embed failed item_id=%s: %s", p.get("item_id"), e)
            self._style(ws, len(products))
            if not self.settings.embed_images: ws.column_dimensions["E"].hidden = True
        if not wb.sheetnames:
            ws = wb.create_sheet("无数据"); ws.append(HEADERS); self._style(ws, 0)
        wb.save(output)
        if self.settings.embed_images and not self.settings.keep_image_cache:
            for path in cleanup: shutil.rmtree(path, ignore_errors=True)
        logging.info("Excel saved: %s", output)
        return output
