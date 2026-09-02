from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import shutil
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import copy
from pathlib import Path
from typing import Any
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage


USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/150.0.0.0 Safari/537.36"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Download Wayfair main images and embed them into split XLSX files."
    )
    parser.add_argument("--input", required=True, help="Source wayfair_products.xlsx")
    parser.add_argument("--output-dir", default="", help="Output directory")
    parser.add_argument("--sheet", default="UniqueProducts")
    parser.add_argument("--chunk-size", type=int, default=300)
    parser.add_argument("--max-workers", type=int, default=6)
    parser.add_argument("--image-size", type=int, default=110)
    parser.add_argument("--cache-dir", default="")
    parser.add_argument("--force-redownload", action="store_true")
    parser.add_argument(
        "--categorized",
        action="store_true",
        help="Treat input as the categorized workbook and embed images per-category sheet.",
    )
    parser.add_argument(
        "--merge-one",
        action="store_true",
        help="With --categorized, write every category as a sheet in ONE workbook instead of many files.",
    )
    return parser.parse_args()


def safe_filename(url: str) -> str:
    return hashlib.sha1(url.encode("utf-8")).hexdigest() + ".jpg"


def download_with_urllib(url: str, timeout: int = 35) -> bytes:
    request = Request(
        url,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "image/avif,image/webp,image/apng,image/*,*/*;q=0.8",
            "Referer": "https://www.wayfair.com/",
        },
    )
    with urlopen(request, timeout=timeout) as response:
        return response.read()


def download_with_curl(url: str, output_path: Path, timeout: int = 45) -> bool:
    curl = shutil.which("curl.exe") or shutil.which("curl")
    if not curl:
        return False

    command = [
        curl,
        "-L",
        "--compressed",
        "--fail",
        "--silent",
        "--show-error",
        "--retry",
        "2",
        "--connect-timeout",
        "15",
        "--max-time",
        str(timeout),
        "-A",
        USER_AGENT,
        "-e",
        "https://www.wayfair.com/",
        "-o",
        str(output_path),
        url,
    ]
    result = subprocess.run(command, capture_output=True, text=True)
    return result.returncode == 0 and output_path.exists() and output_path.stat().st_size > 1000


def convert_to_thumbnail(data: bytes, output_path: Path, image_size: int) -> None:
    with PILImage.open(io.BytesIO(data)) as image:
        image = image.convert("RGB")
        image.thumbnail((image_size, image_size), PILImage.Resampling.LANCZOS)

        canvas = PILImage.new("RGB", (image_size, image_size), "white")
        x = (image_size - image.width) // 2
        y = (image_size - image.height) // 2
        canvas.paste(image, (x, y))
        canvas.save(output_path, "JPEG", quality=78, optimize=True)


def ensure_thumbnail(
    url: str,
    cache_dir: Path,
    image_size: int,
    force_redownload: bool,
) -> tuple[str, Path | None, str | None]:
    cache_path = cache_dir / safe_filename(url)
    if cache_path.exists() and cache_path.stat().st_size > 1000 and not force_redownload:
        return url, cache_path, None

    temp_path = cache_path.with_suffix(".download")
    temp_path.unlink(missing_ok=True)

    last_error = ""
    for attempt in range(1, 4):
        try:
            raw = download_with_urllib(url)
            convert_to_thumbnail(raw, cache_path, image_size)
            return url, cache_path, None
        except Exception as exc:
            last_error = str(exc)

        try:
            if download_with_curl(url, temp_path):
                raw = temp_path.read_bytes()
                convert_to_thumbnail(raw, cache_path, image_size)
                temp_path.unlink(missing_ok=True)
                return url, cache_path, None
        except Exception as exc:
            last_error = str(exc)

        temp_path.unlink(missing_ok=True)
        time.sleep(attempt * 1.5)

    return url, None, last_error or "download failed"


def find_header_index(headers: list[Any], name: str) -> int:
    for index, value in enumerate(headers, start=1):
        if str(value).strip() == name:
            return index
    raise ValueError(f"Header not found: {name}")


def clone_cell_value_and_format(source_cell, target_cell) -> None:
    target_cell.value = source_cell.value
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)

    if source_cell.hyperlink:
        target_cell.hyperlink = copy(source_cell.hyperlink)
        target_cell.style = "Hyperlink"


def style_new_header(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_part(
    source_sheet,
    headers: list[Any],
    rows: list[tuple[Any, ...]],
    start_number: int,
    part_number: int,
    output_dir: Path,
    image_map: dict[str, Path | None],
    image_url_col: int,
    title_col: int,
    image_size: int,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ProductsWithImages"

    sheet.append(["主图"] + headers)
    style_new_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers) + 1)}{len(rows) + 1}"
    sheet.column_dimensions["A"].width = 17

    for source_col in range(1, len(headers) + 1):
        letter = get_column_letter(source_col + 1)
        original_letter = get_column_letter(source_col)
        original_width = source_sheet.column_dimensions[original_letter].width
        sheet.column_dimensions[letter].width = original_width or 14

    embedded = 0
    for local_index, values in enumerate(rows, start=2):
        source_row_number = start_number + local_index - 2

        for source_col, value in enumerate(values, start=1):
            source_cell = source_sheet.cell(source_row_number, source_col)
            target_cell = sheet.cell(local_index, source_col + 1)
            clone_cell_value_and_format(source_cell, target_cell)

        title = values[title_col - 1] if title_col <= len(values) else ""
        image_url = values[image_url_col - 1] if image_url_col <= len(values) else None
        image_path = image_map.get(str(image_url)) if image_url else None

        sheet.row_dimensions[local_index].height = 88
        sheet.cell(local_index, 1).alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        if image_path and image_path.exists():
            image = XLImage(str(image_path))
            image.width = image_size
            image.height = image_size
            image.anchor = f"A{local_index}"
            sheet.add_image(image)
            embedded += 1
        else:
            sheet.cell(local_index, 1).value = "图片下载失败"

        # Keep text readable.
        sheet.cell(local_index, title_col + 1).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    output_path = output_dir / f"wayfair_products_with_images_part{part_number:03d}.xlsx"
    workbook.save(output_path)
    print(
        f"Generated: {output_path} | rows={len(rows)} | embedded={embedded}"
    )
    return output_path


def build_categorized_sheet(
    source_sheet,
    headers: list[Any],
    rows: list[tuple[Any, ...]],
    output_dir: Path,
    image_map: dict[str, Path | None],
    image_url_col: int,
    title_col: int,
    image_size: int,
    category_name: str,
    target_workbook=None,
    target_sheet=None,
):
    safe_name = re.sub(r'[\\/:*?"<>|]+', "_", str(category_name)).strip() or "Category"
    sheet_title = (safe_name or "Category")[:31]
    output_path = output_dir / f"{safe_name}_with_images.xlsx"

    single_file_mode = target_workbook is not None and target_sheet is not None
    if single_file_mode:
        workbook = target_workbook
        sheet = target_sheet
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_title

    sheet.append(["主图"] + headers)
    style_new_header(sheet)
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers) + 1)}{len(rows) + 1}"
    sheet.column_dimensions["A"].width = 17

    for source_col in range(1, len(headers) + 1):
        letter = get_column_letter(source_col + 1)
        original_letter = get_column_letter(source_col)
        original_width = source_sheet.column_dimensions[original_letter].width
        sheet.column_dimensions[letter].width = original_width or 14

    embedded = 0
    for local_index, values in enumerate(rows, start=2):
        source_row = local_index
        for source_col, value in enumerate(values, start=1):
            source_cell = source_sheet.cell(source_row, source_col)
            target_cell = sheet.cell(local_index, source_col + 1)
            clone_cell_value_and_format(source_cell, target_cell)

        title = values[title_col - 1] if title_col <= len(values) else ""
        image_url = values[image_url_col - 1] if image_url_col <= len(values) else None
        image_path = image_map.get(str(image_url)) if image_url else None

        sheet.row_dimensions[local_index].height = 88
        sheet.cell(local_index, 1).alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        if image_path and image_path.exists():
            image = XLImage(str(image_path))
            image.width = image_size
            image.height = image_size
            image.anchor = f"A{local_index}"
            sheet.add_image(image)
            embedded += 1
        else:
            sheet.cell(local_index, 1).value = "图片下载失败"

        if title_col <= len(headers):
            sheet.cell(local_index, title_col + 1).alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    if not single_file_mode:
        workbook.save(output_path)
        print(f"Generated: {output_path} | category={safe_name} | rows={len(rows)} | embedded={embedded}")
    return output_path


def main() -> int:
    args = parse_args()

    input_path = Path(args.input).resolve()
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    output_dir = (
        Path(args.output_dir).resolve()
        if args.output_dir
        else (input_path.parent / ("with_images_categorized" if args.categorized else "with_images"))
    )
    output_dir.mkdir(parents=True, exist_ok=True)

    cache_dir = (
        Path(args.cache_dir).resolve()
        if args.cache_dir
        else output_dir / "image_cache"
    )
    cache_dir.mkdir(parents=True, exist_ok=True)

    if args.categorized:
        failures: list[dict[str, str]] = []
        output_files: list[Path] = []
        source_wb = load_workbook(input_path, data_only=False)

        if args.merge_one:
            merged_wb = Workbook()
            merged_wb.remove(merged_wb.active)
            merged_path = output_dir / "wayfair_categorized_with_images.xlsx"

        for sheet_name in source_wb.sheetnames:
            source_sheet = source_wb[sheet_name]
            headers = [cell.value for cell in source_sheet[1]]
            if "主图URL" not in headers or "商品标题" not in headers:
                print(f"[skip] sheet '{sheet_name}': not a product sheet, skipped.")
                continue

            image_url_col = find_header_index(headers, "主图URL")
            title_col = find_header_index(headers, "商品标题")
            rows = [
                tuple(cell.value for cell in row)
                for row in source_sheet.iter_rows(min_row=2)
                if any(cell.value is not None for cell in row)
            ]
            image_urls = sorted(
                {
                    str(row[image_url_col - 1]).strip()
                    for row in rows
                    if row[image_url_col - 1]
                }
            )
            print(f"[category] {sheet_name}: rows={len(rows)} images={len(image_urls)}")

            image_map: dict[str, Path | None] = {}
            with ThreadPoolExecutor(max_workers=max(1, args.max_workers)) as executor:
                futures = {
                    executor.submit(
                        ensure_thumbnail,
                        url,
                        cache_dir,
                        args.image_size,
                        args.force_redownload,
                    ): url
                    for url in image_urls
                }
                for future in as_completed(futures):
                    url, path, error = future.result()
                    image_map[url] = path
                    if error:
                        failures.append({"url": url, "error": error})

            if args.merge_one:
                target_sheet = merged_wb.create_sheet(title=(re.sub(r'[\\/:*?"<>|]+', "_", sheet_name).strip() or "Category")[:31])
                build_categorized_sheet(
                    source_sheet=source_sheet,
                    headers=headers,
                    rows=rows,
                    output_dir=output_dir,
                    image_map=image_map,
                    image_url_col=image_url_col,
                    title_col=title_col,
                    image_size=args.image_size,
                    category_name=sheet_name,
                    target_workbook=merged_wb,
                    target_sheet=target_sheet,
                )
            else:
                output_files.append(
                    build_categorized_sheet(
                        source_sheet=source_sheet,
                        headers=headers,
                        rows=rows,
                        output_dir=output_dir,
                        image_map=image_map,
                        image_url_col=image_url_col,
                        title_col=title_col,
                        image_size=args.image_size,
                        category_name=sheet_name,
                    )
                )

        if args.merge_one:
            merged_wb.save(merged_path)
            output_files = [merged_path]
            print(f"Generated (merged): {merged_path}")

        failure_path = output_dir / "image_download_failures.csv"
        with failure_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["url", "error"])
            writer.writeheader()
            writer.writerows(failures)

        print("")
        print(f"Done (categorized). Output directory: {output_dir}")
        print(f"Category files created: {len(output_files)}")
        print(f"Image failures: {len(failures)}")
        return 0

    source_wb = load_workbook(input_path, data_only=False)
    if args.sheet not in source_wb.sheetnames:
        raise ValueError(
            f"Sheet '{args.sheet}' not found. Available: {source_wb.sheetnames}"
        )

    source_sheet = source_wb[args.sheet]
    headers = [cell.value for cell in source_sheet[1]]
    image_url_col = find_header_index(headers, "主图URL")
    title_col = find_header_index(headers, "商品标题")

    rows = [
        tuple(cell.value for cell in row)
        for row in source_sheet.iter_rows(min_row=2)
        if any(cell.value is not None for cell in row)
    ]

    image_urls = sorted(
        {
            str(row[image_url_col - 1]).strip()
            for row in rows
            if row[image_url_col - 1]
        }
    )

    print(f"Source rows: {len(rows)}")
    print(f"Unique image URLs: {len(image_urls)}")
    print(f"Cache: {cache_dir}")

    image_map: dict[str, Path | None] = {}
    failures: list[dict[str, str]] = []

    with ThreadPoolExecutor(max_workers=max(1, args.max_workers)) as executor:
        futures = {
            executor.submit(
                ensure_thumbnail,
                url,
                cache_dir,
                args.image_size,
                args.force_redownload,
            ): url
            for url in image_urls
        }

        completed = 0
        for future in as_completed(futures):
            url, path, error = future.result()
            image_map[url] = path
            completed += 1

            if error:
                failures.append({"url": url, "error": error})

            if completed % 25 == 0 or completed == len(image_urls):
                print(f"Downloaded: {completed}/{len(image_urls)}")

    chunk_size = max(50, args.chunk_size)
    output_files: list[Path] = []

    for part_number, start in enumerate(range(0, len(rows), chunk_size), start=1):
        chunk = rows[start : start + chunk_size]
        output_files.append(
            build_part(
                source_sheet=source_sheet,
                headers=headers,
                rows=chunk,
                start_number=start + 2,
                part_number=part_number,
                output_dir=output_dir,
                image_map=image_map,
                image_url_col=image_url_col,
                title_col=title_col,
                image_size=args.image_size,
            )
        )

    failure_path = output_dir / "image_download_failures.csv"
    with failure_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["url", "error"])
        writer.writeheader()
        writer.writerows(failures)

    print("")
    print(f"Done. Output directory: {output_dir}")
    print(f"Parts created: {len(output_files)}")
    print(f"Image failures: {len(failures)}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("Stopped by user.")
        raise SystemExit(130)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
